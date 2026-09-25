import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os


# ============================================================
# COLLEGE ADMISSION MANAGEMENT SYSTEM
# Python + Tkinter + Excel
# ============================================================

FILE_NAME = "college_admissions.xlsx"

USERNAME = "apurva"
PASSWORD = "apurva"


# ============================================================
# EXCEL FUNCTIONS
# ============================================================

def create_excel_file():
    """Create Excel file if it does not exist."""

    if not os.path.exists(FILE_NAME):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Admissions"

        sheet.append([
            "Application ID",
            "Student Name",
            "Gender",
            "Date of Birth",
            "Mobile",
            "Email",
            "Course",
            "12th Percentage",
            "Address"
        ])

        workbook.save(FILE_NAME)


def get_all_records():
    """Read all records from Excel."""

    create_excel_file()

    workbook = load_workbook(FILE_NAME)
    sheet = workbook["Admissions"]

    records = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        records.append(row)

    workbook.close()

    return records


# ============================================================
# MAIN APPLICATION
# ============================================================

class CollegeManagementSystem:

    def __init__(self, root):

        self.root = root
        self.root.title("College Admission Management System")
        self.root.geometry("1100x650")
        self.root.resizable(False, False)

        create_excel_file()

        self.show_login()


    # ========================================================
    # CLEAR WINDOW
    # ========================================================

    def clear_window(self):

        for widget in self.root.winfo_children():
            widget.destroy()


    # ========================================================
    # LOGIN PAGE
    # ========================================================

    def show_login(self):

        self.clear_window()

        self.root.configure(bg="#e8f0fe")

        frame = tk.Frame(
            self.root,
            bg="white",
            width=450,
            height=450
        )

        frame.place(relx=0.5, rely=0.5, anchor="center")

        title = tk.Label(
            frame,
            text="COLLEGE ADMISSION",
            font=("Arial", 24, "bold"),
            bg="white",
            fg="#1a237e"
        )

        title.pack(pady=(45, 5))

        subtitle = tk.Label(
            frame,
            text="Management System",
            font=("Arial", 16),
            bg="white",
            fg="#555555"
        )

        subtitle.pack(pady=(0, 30))


        # Username

        tk.Label(
            frame,
            text="Username",
            font=("Arial", 12, "bold"),
            bg="white"
        ).pack(anchor="w", padx=70)

        self.username_entry = tk.Entry(
            frame,
            font=("Arial", 13),
            width=30
        )

        self.username_entry.pack(pady=(5, 15))


        # Password

        tk.Label(
            frame,
            text="Password",
            font=("Arial", 12, "bold"),
            bg="white"
        ).pack(anchor="w", padx=70)

        self.password_entry = tk.Entry(
            frame,
            font=("Arial", 13),
            width=30,
            show="*"
        )

        self.password_entry.pack(pady=(5, 25))


        # Login Button

        tk.Button(
            frame,
            text="LOGIN",
            font=("Arial", 13, "bold"),
            bg="#3949ab",
            fg="white",
            width=20,
            height=2,
            command=self.login
        ).pack()


        tk.Label(
            frame,
            text="Username: apurva   Password: apurva",
            font=("Arial", 10),
            bg="white",
            fg="gray"
        ).pack(pady=20)


    # ========================================================
    # LOGIN VALIDATION
    # ========================================================

    def login(self):

        username = self.username_entry.get()
        password = self.password_entry.get()

        if username == USERNAME and password == PASSWORD:

            messagebox.showinfo(
                "Login Successful",
                "Welcome to College Admission Management System!"
            )

            self.show_dashboard()

        else:

            messagebox.showerror(
                "Login Failed",
                "Invalid username or password!"
            )


    # ========================================================
    # DASHBOARD
    # ========================================================

    def show_dashboard(self):

        self.clear_window()

        self.root.configure(bg="#f4f6f8")


        # Header

        header = tk.Frame(
            self.root,
            bg="#1a237e",
            height=80
        )

        header.pack(fill="x")

        tk.Label(
            header,
            text="COLLEGE ADMISSION MANAGEMENT SYSTEM",
            font=("Arial", 22, "bold"),
            bg="#1a237e",
            fg="white"
        ).pack(side="left", padx=25, pady=20)


        # Logout

        tk.Button(
            header,
            text="Logout",
            font=("Arial", 11, "bold"),
            bg="#d32f2f",
            fg="white",
            width=10,
            command=self.logout
        ).pack(side="right", padx=20)


        # Navigation

        navigation = tk.Frame(
            self.root,
            bg="white",
            height=70
        )

        navigation.pack(fill="x")


        tk.Button(
            navigation,
            text="Add Record",
            font=("Arial", 11, "bold"),
            width=15,
            command=self.add_record_page
        ).pack(side="left", padx=15, pady=15)


        tk.Button(
            navigation,
            text="View Records",
            font=("Arial", 11, "bold"),
            width=15,
            command=self.view_records_page
        ).pack(side="left", padx=15, pady=15)


        tk.Button(
            navigation,
            text="Search",
            font=("Arial", 11, "bold"),
            width=15,
            command=self.view_records_page
        ).pack(side="left", padx=15, pady=15)


        # Dashboard message

        center = tk.Frame(
            self.root,
            bg="#f4f6f8"
        )

        center.pack(expand=True)


        tk.Label(
            center,
            text="Welcome to College Admission Management System",
            font=("Arial", 22, "bold"),
            bg="#f4f6f8",
            fg="#1a237e"
        ).pack(pady=20)


        tk.Label(
            center,
            text="Manage student admission records easily",
            font=("Arial", 15),
            bg="#f4f6f8",
            fg="#555555"
        ).pack()


        tk.Label(
            center,
            text="\nAdd • View • Search • Update • Delete",
            font=("Arial", 14),
            bg="#f4f6f8",
            fg="#333333"
        ).pack()


    # ========================================================
    # ADD RECORD PAGE
    # ========================================================

    def add_record_page(self):

        self.clear_window()

        self.root.configure(bg="#f4f6f8")


        # Header

        header = tk.Frame(
            self.root,
            bg="#1a237e",
            height=70
        )

        header.pack(fill="x")


        tk.Label(
            header,
            text="ADD STUDENT ADMISSION",
            font=("Arial", 20, "bold"),
            bg="#1a237e",
            fg="white"
        ).pack(side="left", padx=25, pady=18)


        tk.Button(
            header,
            text="Dashboard",
            command=self.show_dashboard
        ).pack(side="right", padx=20)


        # Form

        form = tk.Frame(
            self.root,
            bg="white",
            padx=30,
            pady=20
        )

        form.pack(pady=25)


        # Application ID

        tk.Label(
            form,
            text="Application ID",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=0, column=0, sticky="w", padx=10, pady=8)

        self.app_id = tk.Entry(form, width=30, font=("Arial", 11))
        self.app_id.grid(row=0, column=1, padx=10, pady=8)


        # Student Name

        tk.Label(
            form,
            text="Student Name",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=1, column=0, sticky="w", padx=10, pady=8)

        self.student_name = tk.Entry(form, width=30, font=("Arial", 11))
        self.student_name.grid(row=1, column=1, padx=10, pady=8)


        # Gender

        tk.Label(
            form,
            text="Gender",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=2, column=0, sticky="w", padx=10, pady=8)

        self.gender = ttk.Combobox(
            form,
            values=["Male", "Female", "Other"],
            width=27,
            state="readonly"
        )

        self.gender.grid(row=2, column=1, padx=10, pady=8)


        # DOB

        tk.Label(
            form,
            text="Date of Birth",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=3, column=0, sticky="w", padx=10, pady=8)

        self.dob = tk.Entry(form, width=30, font=("Arial", 11))
        self.dob.grid(row=3, column=1, padx=10, pady=8)


        # Mobile

        tk.Label(
            form,
            text="Mobile",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=4, column=0, sticky="w", padx=10, pady=8)

        self.mobile = tk.Entry(form, width=30, font=("Arial", 11))
        self.mobile.grid(row=4, column=1, padx=10, pady=8)


        # Email

        tk.Label(
            form,
            text="Email",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=5, column=0, sticky="w", padx=10, pady=8)

        self.email = tk.Entry(form, width=30, font=("Arial", 11))
        self.email.grid(row=5, column=1, padx=10, pady=8)


        # Course

        tk.Label(
            form,
            text="Course",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=6, column=0, sticky="w", padx=10, pady=8)

        self.course = ttk.Combobox(
            form,
            values=[
                "Computer Engineering",
                "Information Technology",
                "Mechanical Engineering",
                "Civil Engineering",
                "Electrical Engineering"
            ],
            width=27,
            state="readonly"
        )

        self.course.grid(row=6, column=1, padx=10, pady=8)


        # Percentage

        tk.Label(
            form,
            text="12th Percentage",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=7, column=0, sticky="w", padx=10, pady=8)

        self.percentage = tk.Entry(
            form,
            width=30,
            font=("Arial", 11)
        )

        self.percentage.grid(row=7, column=1, padx=10, pady=8)


        # Address

        tk.Label(
            form,
            text="Address",
            font=("Arial", 11, "bold"),
            bg="white"
        ).grid(row=8, column=0, sticky="nw", padx=10, pady=8)

        self.address = tk.Text(
            form,
            width=28,
            height=3,
            font=("Arial", 11)
        )

        self.address.grid(row=8, column=1, padx=10, pady=8)


        # Buttons

        button_frame = tk.Frame(
            self.root,
            bg="#f4f6f8"
        )

        button_frame.pack()


        tk.Button(
            button_frame,
            text="SAVE RECORD",
            font=("Arial", 11, "bold"),
            bg="#2e7d32",
            fg="white",
            width=15,
            command=self.save_record
        ).grid(row=0, column=0, padx=10)


        tk.Button(
            button_frame,
            text="CLEAR",
            font=("Arial", 11, "bold"),
            bg="#757575",
            fg="white",
            width=15,
            command=self.clear_form
        ).grid(row=0, column=1, padx=10)


    # ========================================================
    # SAVE RECORD
    # ========================================================

    def save_record(self):

        application_id = self.app_id.get().strip()
        name = self.student_name.get().strip()
        gender = self.gender.get()
        dob = self.dob.get().strip()
        mobile = self.mobile.get().strip()
        email = self.email.get().strip()
        course = self.course.get()
        percentage = self.percentage.get().strip()
        address = self.address.get("1.0", tk.END).strip()


        if not application_id or not name or not gender or not course:

            messagebox.showwarning(
                "Missing Data",
                "Please fill all important fields."
            )

            return


        # Check duplicate ID

        records = get_all_records()

        for record in records:

            if str(record[0]) == application_id:

                messagebox.showerror(
                    "Duplicate ID",
                    "Application ID already exists."
                )

                return


        try:

            percentage_value = float(percentage)

            if percentage_value < 0 or percentage_value > 100:

                raise ValueError

        except ValueError:

            messagebox.showerror(
                "Invalid Percentage",
                "Enter percentage between 0 and 100."
            )

            return


        workbook = load_workbook(FILE_NAME)
        sheet = workbook["Admissions"]


        sheet.append([
            application_id,
            name,
            gender,
            dob,
            mobile,
            email,
            course,
            percentage_value,
            address
        ])


        workbook.save(FILE_NAME)
        workbook.close()


        messagebox.showinfo(
            "Success",
            "Student admission record saved successfully!"
        )


        self.clear_form()


    # ========================================================
    # CLEAR FORM
    # ========================================================

    def clear_form(self):

        self.app_id.delete(0, tk.END)
        self.student_name.delete(0, tk.END)

        self.gender.set("")

        self.dob.delete(0, tk.END)
        self.mobile.delete(0, tk.END)
        self.email.delete(0, tk.END)

        self.course.set("")

        self.percentage.delete(0, tk.END)

        self.address.delete("1.0", tk.END)


    # ========================================================
    # VIEW RECORDS
    # ========================================================

    def view_records_page(self):

        self.clear_window()

        self.root.configure(bg="#f4f6f8")


        # Header

        header = tk.Frame(
            self.root,
            bg="#1a237e",
            height=70
        )

        header.pack(fill="x")


        tk.Label(
            header,
            text="STUDENT ADMISSION RECORDS",
            font=("Arial", 20, "bold"),
            bg="#1a237e",
            fg="white"
        ).pack(side="left", padx=25, pady=18)


        tk.Button(
            header,
            text="Dashboard",
            command=self.show_dashboard
        ).pack(side="right", padx=20)


        # Search

        search_frame = tk.Frame(
            self.root,
            bg="#f4f6f8"
        )

        search_frame.pack(pady=15)


        tk.Label(
            search_frame,
            text="Search:",
            font=("Arial", 12, "bold"),
            bg="#f4f6f8"
        ).pack(side="left", padx=5)


        self.search_entry = tk.Entry(
            search_frame,
            width=35,
            font=("Arial", 11)
        )

        self.search_entry.pack(side="left", padx=5)


        tk.Button(
            search_frame,
            text="SEARCH",
            width=12,
            command=self.search_records
        ).pack(side="left", padx=5)


        tk.Button(
            search_frame,
            text="SHOW ALL",
            width=12,
            command=self.load_records
        ).pack(side="left", padx=5)


        # Table

        table_frame = tk.Frame(
            self.root,
            bg="white"
        )

        table_frame.pack(
            padx=15,
            pady=5,
            fill="both",
            expand=True
        )


        columns = (
            "ID",
            "Name",
            "Gender",
            "DOB",
            "Mobile",
            "Email",
            "Course",
            "Percentage",
            "Address"
        )


        self.tree = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings",
            height=18
        )


        for column in columns:

            self.tree.heading(
                column,
                text=column
            )


        widths = {
            "ID": 80,
            "Name": 130,
            "Gender": 80,
            "DOB": 100,
            "Mobile": 110,
            "Email": 160,
            "Course": 160,
            "Percentage": 100,
            "Address": 180
        }


        for column, width in widths.items():

            self.tree.column(
                column,
                width=width,
                anchor="center"
            )


        vertical_scroll = ttk.Scrollbar(
            table_frame,
            orient="vertical",
            command=self.tree.yview
        )

        horizontal_scroll = ttk.Scrollbar(
            table_frame,
            orient="horizontal",
            command=self.tree.xview
        )


        self.tree.configure(
            yscrollcommand=vertical_scroll.set,
            xscrollcommand=horizontal_scroll.set
        )


        self.tree.grid(
            row=0,
            column=0,
            sticky="nsew"
        )

        vertical_scroll.grid(
            row=0,
            column=1,
            sticky="ns"
        )

        horizontal_scroll.grid(
            row=1,
            column=0,
            sticky="ew"
        )


        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)


        # Bottom buttons

        bottom = tk.Frame(
            self.root,
            bg="#f4f6f8"
        )

        bottom.pack(pady=10)


        tk.Button(
            bottom,
            text="UPDATE",
            font=("Arial", 11, "bold"),
            bg="#1976d2",
            fg="white",
            width=15,
            command=self.update_record
        ).pack(side="left", padx=10)


        tk.Button(
            bottom,
            text="DELETE",
            font=("Arial", 11, "bold"),
            bg="#d32f2f",
            fg="white",
            width=15,
            command=self.delete_record
        ).pack(side="left", padx=10)


        tk.Button(
            bottom,
            text="REFRESH",
            font=("Arial", 11, "bold"),
            width=15,
            command=self.load_records
        ).pack(side="left", padx=10)


        self.load_records()


    # ========================================================
    # LOAD RECORDS
    # ========================================================

    def load_records(self):

        if not hasattr(self, "tree"):
            return


        for item in self.tree.get_children():

            self.tree.delete(item)


        records = get_all_records()


        for record in records:

            self.tree.insert(
                "",
                tk.END,
                values=record
            )


    # ========================================================
    # SEARCH
    # ========================================================

    def search_records(self):

        search_text = self.search_entry.get().strip().lower()


        for item in self.tree.get_children():

            self.tree.delete(item)


        records = get_all_records()


        for record in records:

            record_text = " ".join(
                str(value).lower()
                for value in record
                if value is not None
            )


            if search_text in record_text:

                self.tree.insert(
                    "",
                    tk.END,
                    values=record
                )


    # ========================================================
    # UPDATE RECORD
    # ========================================================

    def update_record(self):

        selected = self.tree.selection()


        if not selected:

            messagebox.showwarning(
                "Select Record",
                "Please select a record to update."
            )

            return


        values = self.tree.item(
            selected[0],
            "values"
        )


        self.open_update_window(values)


    # ========================================================
    # UPDATE WINDOW
    # ========================================================

    def open_update_window(self, values):

        window = tk.Toplevel(self.root)

        window.title("Update Student Record")

        window.geometry("500x600")

        window.resizable(False, False)


        tk.Label(
            window,
            text="UPDATE ADMISSION RECORD",
            font=("Arial", 18, "bold")
        ).pack(pady=20)


        frame = tk.Frame(window)

        frame.pack()


        labels = [
            "Application ID",
            "Student Name",
            "Gender",
            "Date of Birth",
            "Mobile",
            "Email",
            "Course",
            "12th Percentage",
            "Address"
        ]


        entries = []


        for i, label in enumerate(labels):

            tk.Label(
                frame,
                text=label,
                font=("Arial", 10, "bold")
            ).grid(
                row=i,
                column=0,
                sticky="w",
                padx=10,
                pady=7
            )


            if label == "Gender":

                entry = ttk.Combobox(
                    frame,
                    values=["Male", "Female", "Other"],
                    width=27,
                    state="readonly"
                )

                entry.set(values[i])

            elif label == "Course":

                entry = ttk.Combobox(
                    frame,
                    values=[
                        "Computer Engineering",
                        "Information Technology",
                        "Mechanical Engineering",
                        "Civil Engineering",
                        "Electrical Engineering"
                    ],
                    width=27,
                    state="readonly"
                )

                entry.set(values[i])

            elif label == "Address":

                entry = tk.Text(
                    frame,
                    width=29,
                    height=3
                )

                entry.insert(
                    "1.0",
                    values[i]
                )

            else:

                entry = tk.Entry(
                    frame,
                    width=30
                )

                entry.insert(
                    0,
                    values[i]
                )


            entry.grid(
                row=i,
                column=1,
                padx=10,
                pady=7
            )


            entries.append(entry)


        def save_update():

            new_values = []


            for i, entry in enumerate(entries):

                if labels[i] == "Address":

                    value = entry.get(
                        "1.0",
                        tk.END
                    ).strip()

                else:

                    value = entry.get().strip()


                new_values.append(value)


            if not new_values[0] or not new_values[1]:

                messagebox.showwarning(
                    "Missing Data",
                    "Application ID and Student Name are required."
                )

                return


            try:

                percentage = float(new_values[7])

                if percentage < 0 or percentage > 100:

                    raise ValueError

            except ValueError:

                messagebox.showerror(
                    "Invalid Percentage",
                    "Enter percentage between 0 and 100."
                )

                return


            workbook = load_workbook(FILE_NAME)

            sheet = workbook["Admissions"]


            for row in sheet.iter_rows(
                min_row=2
            ):

                if str(row[0].value) == str(values[0]):

                    for column_number in range(
                        1,
                        10
                    ):

                        row[column_number - 1].value = new_values[
                            column_number - 1
                        ]


                    break


            workbook.save(FILE_NAME)

            workbook.close()


            messagebox.showinfo(
                "Updated",
                "Record updated successfully!"
            )


            window.destroy()

            self.load_records()


        tk.Button(
            window,
            text="UPDATE RECORD",
            font=("Arial", 11, "bold"),
            bg="#1976d2",
            fg="white",
            width=20,
            command=save_update
        ).pack(pady=20)


    # ========================================================
    # DELETE RECORD
    # ========================================================

    def delete_record(self):

        selected = self.tree.selection()


        if not selected:

            messagebox.showwarning(
                "Select Record",
                "Please select a record to delete."
            )

            return


        values = self.tree.item(
            selected[0],
            "values"
        )


        application_id = values[0]


        confirm = messagebox.askyesno(
            "Confirm Delete",
            "Are you sure you want to delete this record?"
        )


        if not confirm:

            return


        workbook = load_workbook(FILE_NAME)

        sheet = workbook["Admissions"]


        for row in range(
            2,
            sheet.max_row + 1
        ):

            if str(
                sheet.cell(row, 1).value
            ) == str(application_id):

                sheet.delete_rows(
                    row,
                    1
                )

                break


        workbook.save(FILE_NAME)

        workbook.close()


        messagebox.showinfo(
            "Deleted",
            "Record deleted successfully!"
        )


        self.load_records()


    # ========================================================
    # LOGOUT
    # ========================================================

    def logout(self):

        confirm = messagebox.askyesno(
            "Logout",
            "Are you sure you want to logout?"
        )


        if confirm:

            self.show_login()


# ============================================================
# START APPLICATION
# ============================================================

if __name__ == "__main__":

    root = tk.Tk()

    app = CollegeManagementSystem(root)

    root.mainloop()